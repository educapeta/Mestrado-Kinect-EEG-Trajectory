# -*- coding: utf-8 -*-
"""Valida o .docx gerado: integridade do zip e leitura pelo python-docx."""

import zipfile
from pathlib import Path

import docx

import _fill_report as F

path = Path(F.DST)
print("existe:", path.exists(), "bytes:", path.stat().st_size)

with zipfile.ZipFile(path) as zf:
    print("zip testzip (None = ok):", zf.testzip())

document = docx.Document(str(path))
print("paragrafos:", len(document.paragraphs), "tabelas:", len(document.tables))

tabelas_ok = all(len(t.rows) > 0 for t in document.tables)
print("todas as tabelas com linhas:", tabelas_ok)

celulas_vazias = []
for ti, table in enumerate(document.tables):
    if ti in (3, 4, 5, 6, 7):
        txt = table.rows[0].cells[0].text.strip()
        celulas_vazias.append((ti, len(txt)))
print("blocos de texto (indice, chars):", celulas_vazias)
