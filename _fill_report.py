# -*- coding: utf-8 -*-
"""Preenche o template do Relatorio Tecnico de Bolsas (PPGNeuro/UFSC), 2o semestre.

Uso:  .venv\\Scripts\\python.exe _fill_report.py
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

import _report_content as C

CB_OFF = "\u2610"  # caixa vazia
CB_ON = "\u2612"   # caixa marcada

SRC = Path(
    r"C:\Users\elgut\OneDrive\Desktop\Mestrado_Projeto\Bolsa Fapesc\Relatórios\2-Semestre"
    r"\RELATÓRIO TÉCNICO DE BOLSAS ACADÊMICAS PPGNeuro_revisado .docx"
)
DST = Path(
    r"C:\Users\elgut\OneDrive\Desktop\Mestrado_Projeto\Bolsa Fapesc\Relatórios\2-Semestre"
    r"\relatorio_tecnico_2semestre_Eduardo_Larsen_Guths.docx"
)

CREDITOS_EXIGIDOS = "24"
CREDITOS_CONCLUIDOS = "24"
DATA_INGRESSO = "01/09/2025"
DATA_QUALIFICACAO = "26/03/2026"
PREVISAO_DEFESA = "31/07/2027"
LABORATORIO = "PsiCow Lab (Neuroscience Coworking Lab) – PPGNeuro/UFSC"
AREA = "Neurociências"
LINHA = "Eletrofisiologia / Interfaces Cérebro-Máquina"
CARGA_HORARIA = "40"


def iter_blocks(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def set_text(paragraph, text: str) -> None:
    """Troca o texto do paragrafo preservando a formatacao do 1o run."""
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for extra in runs[1:]:
            extra._element.getparent().remove(extra._element)
        return
    pPr = paragraph._p.find(qn("w:pPr"))
    for child in list(paragraph._p):
        if child is not pPr:
            paragraph._p.remove(child)
    paragraph.add_run(text)


def append_text(cell, text: str) -> None:
    """Adiciona um paragrafo com o texto no fim da celula."""
    paragraph = cell.add_paragraph()
    set_text(paragraph, text)


def unique_cells(row):
    seen = []
    for cell in row.cells:
        if all(cell._tc is not s for s in seen):
            seen.append(cell)
    return seen


def fill_cell(cell, lines) -> None:
    paras = cell.paragraphs
    while len(paras) < len(lines):
        cell.add_paragraph()
        paras = cell.paragraphs
    for i, para in enumerate(paras):
        set_text(para, lines[i] if i < len(lines) else "")


def q(label: str, done: bool) -> str:
    return f"{(CB_ON if done else CB_OFF)} {label}"


# --- edicoes por indice de paragrafo (mesma numeracao do _inspect_docx) ------
P_EDITS = {
    12: f"Total de créditos em disciplinas exigidos pelo programa: {CREDITOS_EXIGIDOS}",
    13: f"Total de créditos já concluídos (com aprovação): {CREDITOS_CONCLUIDOS}",
    14: "Percentual de créditos concluídos:   " + "   ".join([
        q("25%", False), q("50%", False), q("75%", False), q("100%", True),
    ]),
    18: q("Não realizada", False),
    19: q("Agendada para: __/__/____", False),
    20: q(f"Concluída em: {DATA_QUALIFICACAO}", True),
    21: q(f"Previsão de Defesa Final: {PREVISAO_DEFESA}", False),
    25: q("Não realizada", False),
    26: q("Agendada para: __/__/____", False),
    27: q("Concluída em: __/__/____", False),
    28: q(f"Previsão de Defesa Final: {PREVISAO_DEFESA}", False),
    38: q(
        "Sim. Quais? Participação como ouvinte nos seminários do PPG Neurociências / "
        "PsiCow Lab realizados no período (certificados em anexo).",
        True,
    ),
    39: q("Não.", False),
    41: q("Sim.", False),
    42: q("Não", True),
    44: q("Sim. Quais?", False),
    45: q("Não", True),
    49: q("Não (Dedicação exclusiva)", True),
    50: q("Sim.", False),
    52: q("Sim, cumpro a carga horária mínima exigida.", True),
    53: q('Não (Justificar em "Alterações ou Dificuldades")', False),
    54: (
        "Carga horária semanal total dedicada as atividades relacionadas ao "
        f"PPG Neuro : {CARGA_HORARIA} horas."
    ),
    61: q("Revisão Bibliográfica/Sistemática iniciada", True),
    62: q("Revisão Bibliográfica/Sistemática finalizada", True),
    63: q("Submissão no comitê de ética (se aplicável)", True),
    64: q("Aprovação no comitê de ética (se aplicável)", False),
    65: q("Coleta de Dados/Experimentação iniciada", False),
    66: q("Coleta de Dados/Experimentação concluída", False),
    67: q("Análise de Dados", True),
    68: q("Redação do Manuscrito/Dissertação/Tese", True),
    70: "Percentual de pesquisa concluído:   " + "   ".join([
        q("25%", False), q("50%", True), q("75%", False), q("100%", False),
    ]),
    74: q("Execução conforme planejado", False),
    75: q("Execução com atrasos menores (não impacta a defesa)", True),
    76: q("Execução com atrasos significativos (requer revisão do cronograma)", False),
    95: q("Abaixo do esperado", False),
    96: q("Dentro do esperado", True),
    97: q("Acima do esperado", False),
    100: q("Falta de tempo", False),
    101: q("Problemas metodológicos", False),
    102: q("Falta de infraestrutura", True),
    103: q("Dificuldades na análise de dados", False),
    104: q("Escrita científica", False),
    105: q("Organização pessoal", False),
    106: q("Demandas externas (trabalho, família)", False),
    109: q("Sim, com segurança", False),
    110: q("Provável, com ajustes", True),
    111: q("Risco moderado", False),
    112: q("Alto risco", False),
}


def edit_tables(document) -> list[str]:
    log = []
    tables = document.tables

    # ---- T00: modalidade e tipo de relatorio -------------------------------
    t00 = tables[0]
    fill_cell(t00.rows[1].cells[0], [
        q("Mestrado", True), q("Doutorado", False), q("Pós-Doutorado", False),
    ])
    fill_cell(t00.rows[1].cells[-1], [
        q("Mestrado Sanduíche no exterior", False),
        q("Doutorado Sanduíche no exterior", False),
        q("Pós-Doutorado Sanduíche no exterior", False),
    ])
    fill_cell(t00.rows[3].cells[0], [q("Parcial", True), q("Final", False)])
    log.append("T00 modalidade/tipo marcados")

    # ---- T01: dados do bolsista -------------------------------------------
    # As linhas de resposta nao tem celulas proprias: o rotulo esta em uma
    # celula mesclada verticalmente, entao o valor entra como novo paragrafo.
    t01 = tables[1]
    valores = {
        0: "Eduardo Larsen Güths",
        2: "Eduardo Luiz Gasnhar Moreira",
        4: "Matheus Mangini Bertuzzo",
        6: LABORATORIO,
        8: AREA,
        10: LINHA,
        12: C.TITULO,
        14: DATA_INGRESSO,
    }
    for row_idx, valor in valores.items():
        cell = t01.rows[row_idx].cells[0]
        if len(cell.text.strip()) > len(t01.rows[row_idx].cells[0].paragraphs[0].text):
            log.append(f"AVISO: T01 R{row_idx} ja tinha resposta; preservada")
            continue
        append_text(cell, valor)
    log.append("T01 dados do bolsista preenchidos")

    # ---- T02: producao tecnico-cientifica ---------------------------------
    t02 = tables[2]
    producao = {
        1: ("0", "Revista: — (nenhum artigo submetido no período)"),
        2: ("0", "DOI ou Link: —"),
        3: ("2", "Nome do evento: seminários do PPG Neurociências / PsiCow Lab "
                 "(ouvinte, mar. e jun./2026)"),
        4: ("1", f"{CB_ON} Oral {CB_OFF} Poster — Qualificação de Mestrado "
                 f"({DATA_QUALIFICACAO})"),
        5: ("1", "Especificar: documentação ética completa (TCLE v5, projeto "
                 "detalhado, folha de rosto, declaração institucional e modelo de "
                 "anonimização de dados) submetida ao CEP"),
    }
    for row_idx, (qtd, status) in producao.items():
        cells = unique_cells(t02.rows[row_idx])
        fill_cell(cells[1], [qtd])
        fill_cell(cells[2], [status])
    log.append("T02 produção preenchida")

    # ---- T03..T07: blocos de texto ----------------------------------------
    blocos = [
        (3, C.RESUMO_PROPOSTA, "7.1 resumo da proposta"),
        (4, C.ALTERACOES_DIFICULDADES, "alterações/dificuldades"),
        (5, C.RESULTADOS, "resultados alcançados"),
        (6, C.CONCLUSOES, "conclusões"),
        (7, C.PERSPECTIVAS, "perspectivas"),
    ]
    for table_idx, linhas, nome in blocos:
        cell = tables[table_idx].rows[0].cells[0]
        fill_cell(cell, linhas)
        log.append(f"T{table_idx:02d} {nome} preenchido ({len(linhas)} parágrafos)")

    return log


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"template nao encontrado: {SRC}")
    document = docx.Document(str(SRC))

    log = []
    pi = 0
    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            if pi in P_EDITS:
                set_text(block, P_EDITS[pi])
                log.append(f"P{pi:03d} -> {P_EDITS[pi][:60]!r}")
            pi += 1

    log.extend(edit_tables(document))

    DST.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(DST))

    report = [f"OK -> {DST}"] + ["  " + line for line in log]
    Path(r"c:\mestradopy\_fill_log.txt").write_text(
        "\n".join(report), encoding="utf-8"
    )
    print(f"OK -> {DST} ({len(log)} edicoes)")


if __name__ == "__main__":
    main()
